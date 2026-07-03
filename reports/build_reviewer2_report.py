from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path


ROOT = Path("/Users/sudipphuyal/Developments/zkp-with-snarkjs")
OUT = ROOT / "reports" / "reviewer2-smart-contract-verification-report.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    r.font.size = Pt(9)


def style_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    for cell in table.rows[0].cells:
        shade_cell(cell, "E8EEF5")
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)

for name, size, color in [
    ("Heading 1", 15, "2E74B5"),
    ("Heading 2", 12.5, "2E74B5"),
    ("Heading 3", 11.5, "1F4D78"),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Reviewer #2 Response Evidence Report")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor.from_string("0B2545")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("Smart Contract Verification, Static Analysis, and Formal Verification Status")
subtitle_run.font.size = Pt(11)
subtitle_run.italic = True
subtitle_run.font.color.rgb = RGBColor.from_string("555555")

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Repository: zkp-with-snarkjs | Branch context: upgrade-to-vcem | Report date: 2026-07-03").font.size = Pt(9)

doc.add_heading("Executive Answer", level=1)
p = doc.add_paragraph()
p.add_run("Smart-contract-oriented validation exists and was partially executed. ").bold = True
p.add_run(
    "The repository includes executed unit/integration/matrix/audit-verifier tests, Solhint static analysis, coverage, gas reporting, dependency audit, and an independent audit-verifier test suite. "
    "Slither static analysis is configured in npm scripts and GitHub Actions, but it was not executed locally because the Slither binary was unavailable. "
    "No completed formal verification proof currently exists."
)

doc.add_heading("Reviewer-Safe Manuscript Wording", level=1)
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Inches(0.25)
quote.paragraph_format.right_indent = Inches(0.25)
quote.add_run(
    "We performed automated smart-contract validation through unit and integration tests, a 60-case/120-outcome VCEM authorization matrix, an independent audit-verifier test suite, Solhint static analysis, Solidity coverage, gas reporting, and dependency auditing. "
    "Slither analysis is configured in CI but was not executed in the local evidence run due to a missing local Slither binary. "
    "The current artifact does not include formal verification with tools such as Certora, Scribble/Echidna, Foundry invariant proofs, or SMTChecker proofs; these are identified as future hardening work."
).italic = True

doc.add_heading("Verification Status Matrix", level=1)
rows = [
    ["Category", "Tool / command", "Status", "Evidence / result", "How to report"],
    ["Compile verification", "HARDHAT_DISABLE_DOWNLOADS=true npm run compile", "Executed: pass", "Offline Hardhat compile succeeded with locally pinned solc packages.", "Can report as reproducible offline compile."],
    ["Functional tests", "npm test", "Executed: pass", "85 passing tests.", "Can report as functional and integration validation."],
    ["VCEM matrix", "npm run test:vcem:matrix", "Executed: pass", "60 policy cases and 120 outcomes.", "Can report as VCEM authorization matrix evidence."],
    ["Audit verifier tests", "npm run test:audit-verifier", "Executed: pass", "17 passing tamper-detection tests.", "Can report as independent verifier validation."],
    ["FHIR tests", "npm run test:fhir", "Executed: pass", "6 passing fixture lifecycle tests.", "Can report as fixture-based FHIR mapping only."],
    ["Property tests", "npm run test:property", "Executed: pass", "2 canonical-hash property tests.", "Can report as property-based hash checks."],
    ["Static analysis", "npm run solhint", "Executed: pass with warnings", "Solhint exited 0; warnings mostly NatSpec/gas style.", "Can report as Solhint static analysis with warnings."],
    ["Static analysis", "npm run slither", "Configured, not locally executed", "Failed locally: slither: command not found. CI uses crytic/slither-action.", "Do not claim completed Slither findings from local run."],
    ["Coverage", "npm run coverage", "Executed: pass", "87 passing, 4 pending; 77.66% statement coverage.", "Can report coverage evidence."],
    ["Gas report", "npm run gas", "Executed: pass", "6 VCEM tests passing with gas report.", "Can report gas profiling was executed."],
    ["Dependency audit", "npm run security:audit", "Executed: fail", "54 vulnerabilities: 19 low, 21 moderate, 11 high, 3 critical.", "Report as unresolved dependency risk, not clean security sign-off."],
    ["Bytecode verifier", "npm run audit:verify-bytecode", "Implemented, deployment-dependent", "Verifier exists; needs deployed manifest/RPC for live run.", "Do not claim live deployed bytecode verification unless run after Besu deployment."],
    ["Formal verification", "Certora / SMTChecker / Scribble / Echidna / Foundry invariants", "Not implemented as completed proof", "No formal proof artifacts found.", "Must be reported as future work or limitation."],
]
table = doc.add_table(rows=len(rows), cols=len(rows[0]))
for r_idx, row in enumerate(rows):
    for c_idx, text in enumerate(row):
        set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))
style_table(table, [1.2, 1.75, 1.35, 2.2, 1.75])

doc.add_heading("Static Analysis Evidence", level=1)
add_bullets(doc, [
    "Solhint is installed and configured through .solhint.json and package.json script npm run solhint.",
    "Solhint was executed and exited successfully, but emitted warnings. The warning set is primarily documentation and gas-style guidance across retained legacy contracts and VCEM contracts.",
    "Slither is configured through package.json and slither.config.json, and GitHub Actions includes crytic/slither-action@v0.4.1.",
    "Local Slither execution was attempted through npm run slither and failed because the slither binary was not installed in the local environment.",
])

doc.add_heading("Formal Verification Status", level=1)
add_bullets(doc, [
    "No completed formal verification evidence exists in the current repository.",
    "No Certora specification, Scribble annotations, Echidna campaign, Foundry invariant test suite, or SMTChecker proof report was found.",
    "The current property tests are useful validation, but they are not a substitute for formal verification.",
    "The manuscript should not state that formal verification has been completed unless a new formal verification workflow is implemented, executed, and archived.",
])

doc.add_heading("Smart Contract Verification vs. Audit Verification", level=1)
p = doc.add_paragraph()
p.add_run("Important distinction: ").bold = True
p.add_run(
    "the repository contains an independent audit verifier for VCEM compliance from blockchain evidence. "
    "That verifier is not the same as formal smart contract verification. It reconstructs consent/access evidence, checks signatures and consent-state binding, and has a tamper-detection test suite."
)

doc.add_page_break()
doc.add_heading("Evidence Files and Commands", level=1)
rows = [
    ["Evidence item", "Path / command"],
    ["Solhint config", ".solhint.json"],
    ["Slither config", "slither.config.json"],
    ["CI static-analysis hook", ".github/workflows/vcem-ci.yml"],
    ["Coverage report", "coverage/index.html, coverage/lcov.info, coverage.json"],
    ["Gas report command", "npm run gas"],
    ["Audit verifier", "scripts/auditVerify.ts"],
    ["Audit verifier tests", "test/AuditVerifier.ts"],
    ["VCEM contracts", "contracts/vcem/VCEMRegistry.sol, VCEMConsent.sol, VCEMAudit.sol"],
    ["Dependency risk register", "docs/dependency-risk-register.md"],
]
table = doc.add_table(rows=len(rows), cols=2)
for r_idx, row in enumerate(rows):
    for c_idx, text in enumerate(row):
        set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))
style_table(table, [2.25, 5.75])

doc.add_heading("Recommended Reviewer #2 Response", level=1)
p = doc.add_paragraph()
p.add_run(
    "We agree with the reviewer that static analysis and formal verification should be clearly distinguished. "
    "In the revised artifact, static analysis is represented by Solhint execution and Slither CI configuration. "
    "Formal verification is not claimed as completed. We now report the implemented VCEM evidence as automated testing, independent audit-verifier validation, coverage/gas profiling, and configured static-analysis tooling, while identifying full formal verification as future work."
)

doc.add_heading("Limitations to State Explicitly", level=1)
add_bullets(doc, [
    "No completed Slither local result is available from this execution.",
    "No formal verification proof artifacts are available.",
    "Dependency audit currently fails and should be disclosed as unresolved dependency risk.",
    "Bytecode verification is implemented but requires a deployed Besu network, manifest, and RPC endpoint for live execution evidence.",
    "Legacy ZKP tests are experimental and are not evidence of VCEM authorization proof.",
])

doc.add_heading("Conclusion", level=1)
p = doc.add_paragraph()
p.add_run("For the final paper: ").bold = True
p.add_run(
    "state that the implementation was validated with automated contract tests, VCEM authorization matrix tests, audit-verifier tamper tests, Solhint static analysis, coverage, gas profiling, and dependency audit. "
    "Do not state that formal verification was completed. If Reviewer #2 requires formal verification, the next implementation step should add a dedicated formal-verification workflow and archive its outputs."
)

doc.save(OUT)
print(OUT)
